from contextlib import asynccontextmanager
from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
from typing import List, Dict, Optional
import time, uuid, os, json, re, asyncio, aiohttp, smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
from datetime import datetime
import subprocess, shutil, hashlib
from concurrent.futures import ThreadPoolExecutor
from collections import OrderedDict
from threading import Lock
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import qrcode
from fpdf import FPDF
from docx import Document
from sympy.parsing.sympy_parser import (
    parse_expr, standard_transformations, implicit_multiplication_application
)
from deep_translator import GoogleTranslator
from bs4 import BeautifulSoup
from PIL import Image
import logging
import logging.handlers
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded

try:
    import pdfkit
    PDFKIT_AVAILABLE = True
except ImportError:
    PDFKIT_AVAILABLE = False

try:
    import markdown2
    MARKDOWN2_AVAILABLE = True
except ImportError:
    MARKDOWN2_AVAILABLE = False

try:
    import fitz
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import mammoth
    MAMMOTH_AVAILABLE = True
except ImportError:
    MAMMOTH_AVAILABLE = False

# ==================== CONFIGURATION ====================
OUTPUT_DIR  = os.getenv("OUTPUT_DIR",  "/opt/ai-backend/output")
SEARXNG_URL = os.getenv("SEARXNG_URL", "http://localhost:8081")
OLLAMA_URL  = os.getenv("OLLAMA_URL",  "http://localhost:11434")
TOOL_MODEL  = os.getenv("TOOL_MODEL",  "qwen3:4b-instruct")
CHAT_MODEL  = os.getenv("CHAT_MODEL",  "qwen3:14b")
MAX_TOOL_LOOPS = 3  # max tool call iterations per request

TOOL_KEYWORDS = {
    "create", "make", "generate", "calculate", "compute", "translate",
    "list", "qr", "pdf", "docx", "word", "chart", "graph", "resize",
    "merge", "split", "compress", "scrape", "email", "send"
}

EMAIL_CONFIG = {
    "smtp_server": os.getenv("SMTP_SERVER", "smtp.gmail.com"),
    "smtp_port":   int(os.getenv("SMTP_PORT", "587")),
    "username":    os.getenv("EMAIL_USERNAME"),
    "password":    os.getenv("EMAIL_PASSWORD"),
}

# ==================== LOGGING ====================
def setup_logging():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    handler = logging.handlers.RotatingFileHandler(
        os.path.join(OUTPUT_DIR, 'ai_agent.log'),
        maxBytes=10*1024*1024, backupCount=5
    )
    fmt = logging.Formatter('%(asctime)s - %(name)s - %(levelname)s - %(message)s')
    handler.setFormatter(fmt)
    logger = logging.getLogger(__name__)
    logger.setLevel(logging.INFO)
    logger.addHandler(handler)
    console = logging.StreamHandler()
    console.setFormatter(fmt)
    logger.addHandler(console)
    return logger

logger = setup_logging()

N_TOOL_THREADS = max(1, (os.cpu_count() or 4) - 1)
executor = ThreadPoolExecutor(max_workers=N_TOOL_THREADS)

# ==================== UTILITY ====================
def sanitize_filename(filename: str) -> str:
    filename = os.path.basename(filename)
    filename = re.sub(r'[^\w\s\-\.]', '', filename)
    filename = filename.strip('. ')
    if not filename:
        filename = f"file_{uuid.uuid4().hex[:8]}"
    return filename

def detect_language(text: str) -> str:
    if not text:
        return "en"
    tl = text.lower()
    ro_words = ['este','sunt','acest','această','pentru','cum','când','unde','și','sau',
                'dar','că','de','la','în','pe','cu','să','vreau','faci','creează','fișier']
    fr_words = ['est','sont','avec','pour','dans','une','des','les','qui','que',
                'être','créer','faire','voulez','veux','fichier','vous']
    ro = sum(1 for w in ro_words if f' {w} ' in f' {tl} ')
    fr = sum(1 for w in fr_words if f' {w} ' in f' {tl} ')
    if ro >= 2: return "ro"
    if fr >= 2: return "fr"
    return "en"

def check_wkhtmltopdf() -> bool:
    return shutil.which('wkhtmltopdf') is not None

def is_safe_url(url: str) -> bool:
    from urllib.parse import urlparse
    try:
        p = urlparse(url)
        if p.scheme not in ('http', 'https'):
            return False
        host = (p.hostname or '').lower()
        blocked = {'localhost', '127.0.0.1', '0.0.0.0', '::1', '169.254.169.254'}
        blocked_prefixes = ('10.', '192.168.', '172.')
        if host in blocked or any(host.startswith(b) for b in blocked_prefixes):
            return False
        return True
    except Exception:
        return False

def safe_calculate(expression: str) -> str:
    if not re.match(r'^[\d\s\+\-\*\/\(\)\.\^%]+$', expression):
        return "Error: Invalid characters. Only numbers and +-*/()^.% allowed."
    try:
        transformations = standard_transformations + (implicit_multiplication_application,)
        result = parse_expr(expression, transformations=transformations, evaluate=True)
        return f"Result: {result.evalf()}"
    except Exception as e:
        return f"Error: {str(e)}"

def validate_environment():
    if not OUTPUT_DIR:
        raise EnvironmentError("OUTPUT_DIR must be set")
    os.makedirs(OUTPUT_DIR, exist_ok=True)

# ==================== DOCUMENT TOOLS ====================
class DocumentTools:
    @staticmethod
    def markdown_to_pdf(markdown_content: str, filename: str) -> str:
        try:
            filename = sanitize_filename(filename)
            filepath = os.path.join(OUTPUT_DIR, filename)
            if not MARKDOWN2_AVAILABLE:
                return "Error: markdown2 not available"
            html_content = markdown2.markdown(markdown_content)
            html_template = f"""<!DOCTYPE html>
<html><head><meta charset="UTF-8">
<style>body{{font-family:Arial,sans-serif;margin:40px;line-height:1.6}}
h1,h2,h3{{color:#333}}code{{background:#f4f4f4;padding:2px 4px;border-radius:3px}}
pre{{background:#f4f4f4;padding:10px;border-radius:5px}}</style>
</head><body>{html_content}</body></html>"""
            if PDFKIT_AVAILABLE and check_wkhtmltopdf():
                try:
                    pdfkit.from_string(html_template, filepath, {
                        'page-size':'A4','margin-top':'0.75in',
                        'margin-right':'0.75in','margin-bottom':'0.75in',
                        'margin-left':'0.75in','encoding':'UTF-8'})
                    return f"PDF created: {filepath}"
                except Exception as e:
                    logger.warning(f"pdfkit failed: {e}")
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            doc = SimpleDocTemplate(filepath, pagesize=letter)
            styles = getSampleStyleSheet()
            story = []
            for line in markdown_content.split('\n'):
                if line.strip():
                    story.append(Paragraph(line, styles["Normal"]))
                    story.append(Spacer(1, 12))
            doc.build(story)
            return f"PDF created: {filepath}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def markdown_to_html(markdown_content: str, filename: str) -> str:
        try:
            filename = sanitize_filename(filename)
            if not MARKDOWN2_AVAILABLE:
                return "Error: markdown2 not available"
            html_content = markdown2.markdown(markdown_content)
            full_html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>Document</title>
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/github-markdown-css/5.1.0/github-markdown.min.css">
<style>.markdown-body{{box-sizing:border-box;min-width:200px;max-width:980px;margin:0 auto;padding:45px}}</style>
</head><body><article class="markdown-body">{html_content}</article></body></html>"""
            filepath = os.path.join(OUTPUT_DIR, filename)
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(full_html)
            return f"HTML created: {filepath}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def docx_to_pdf(docx_path: str, filename: str) -> str:
        try:
            filename = sanitize_filename(filename)
            filepath = os.path.join(OUTPUT_DIR, filename)
            if shutil.which('libreoffice'):
                cmd = ['libreoffice','--headless','--convert-to','pdf','--outdir',OUTPUT_DIR,docx_path]
                result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
                if result.returncode == 0:
                    base_name = os.path.basename(docx_path).rsplit('.', 1)[0]
                    expected = os.path.join(OUTPUT_DIR, f"{base_name}.pdf")
                    if os.path.exists(expected):
                        os.rename(expected, filepath)
                        return f"PDF created: {filepath}"
            if not MAMMOTH_AVAILABLE:
                return "Error: mammoth not available"
            with open(docx_path, "rb") as f:
                result = mammoth.convert_to_html(f)
                html = result.value
                if PDFKIT_AVAILABLE and check_wkhtmltopdf():
                    pdfkit.from_string(html, filepath, {'page-size':'A4','encoding':'UTF-8'})
                    return f"PDF created: {filepath}"
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=12)
                for line in BeautifulSoup(html, 'html.parser').get_text().split('\n'):
                    if line.strip():
                        pdf.multi_cell(0, 10, line)
                pdf.output(filepath)
                return f"PDF created: {filepath}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def merge_pdfs(pdf_files: List[str], output_filename: str) -> str:
        try:
            from PyPDF2 import PdfMerger
            output_filename = sanitize_filename(output_filename)
            merger = PdfMerger()
            for pdf_file in pdf_files:
                path = pdf_file if os.path.exists(pdf_file) else os.path.join(OUTPUT_DIR, pdf_file)
                if not os.path.exists(path):
                    merger.close()
                    return f"Error: {pdf_file} not found"
                merger.append(path)
            output_path = os.path.join(OUTPUT_DIR, output_filename)
            merger.write(output_path)
            merger.close()
            return f"PDFs merged: {output_path}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def split_pdf(pdf_path: str, pages: List[int] = None, output_prefix: str = "split") -> str:
        try:
            from PyPDF2 import PdfReader, PdfWriter
            output_prefix = sanitize_filename(output_prefix)
            if not os.path.exists(pdf_path):
                pdf_path = os.path.join(OUTPUT_DIR, pdf_path)
            if not os.path.exists(pdf_path):
                return "Error: File not found"
            reader = PdfReader(pdf_path)
            total = len(reader.pages)
            if pages:
                writer = PdfWriter()
                for n in pages:
                    if 1 <= n <= total:
                        writer.add_page(reader.pages[n - 1])
                out = os.path.join(OUTPUT_DIR, f"{output_prefix}_pages.pdf")
                with open(out, 'wb') as f:
                    writer.write(f)
                return f"PDF split: {out}"
            else:
                results = []
                for i in range(total):
                    writer = PdfWriter()
                    writer.add_page(reader.pages[i])
                    out = os.path.join(OUTPUT_DIR, f"{output_prefix}_page_{i+1}.pdf")
                    with open(out, 'wb') as f:
                        writer.write(f)
                    results.append(os.path.basename(out))
                return f"Split into {total} files: {', '.join(results)}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def compress_pdf(input_path: str, output_filename: str) -> str:
        try:
            from PyPDF2 import PdfReader, PdfWriter
            output_filename = sanitize_filename(output_filename)
            if not os.path.exists(input_path):
                input_path = os.path.join(OUTPUT_DIR, input_path)
            if not os.path.exists(input_path):
                return "Error: File not found"
            reader = PdfReader(input_path)
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            output_path = os.path.join(OUTPUT_DIR, output_filename)
            writer.add_metadata(reader.metadata)
            with open(output_path, 'wb') as f:
                writer.write(f)
            orig = os.path.getsize(input_path)
            comp = os.path.getsize(output_path)
            ratio = (1 - comp / orig) * 100 if orig > 0 else 0
            return f"PDF compressed: {output_path} (Reduced by {ratio:.1f}%)"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def extract_images_from_pdf(pdf_path: str, output_dir: str = "extracted_images") -> str:
        try:
            if not PYMUPDF_AVAILABLE:
                return "Error: PyMuPDF not available"
            output_dir = sanitize_filename(output_dir)
            if not os.path.exists(pdf_path):
                pdf_path = os.path.join(OUTPUT_DIR, pdf_path)
            if not os.path.exists(pdf_path):
                return "Error: File not found"
            out_path = os.path.join(OUTPUT_DIR, output_dir)
            os.makedirs(out_path, exist_ok=True)
            doc = fitz.open(pdf_path)
            count = 0
            for pg in range(len(doc)):
                for idx, img in enumerate(doc.load_page(pg).get_images()):
                    base = doc.extract_image(img[0])
                    if base:
                        fn = f"page_{pg+1}_img_{idx+1}.{base['ext']}"
                        with open(os.path.join(out_path, fn), 'wb') as f:
                            f.write(base['image'])
                        count += 1
            doc.close()
            return f"Extracted {count} images to: {out_path}"
        except Exception as e:
            return f"Error: {str(e)}"

# ==================== EMAIL TOOLS ====================
class EmailTools:
    @staticmethod
    def send_email(to_email: str, subject: str, body: str,
                   attachments: List[str] = None, html_body: str = None) -> str:
        try:
            username = EMAIL_CONFIG.get('username')
            password = EMAIL_CONFIG.get('password')
            if not username or not password:
                return "Error: Email credentials not configured (set EMAIL_USERNAME and EMAIL_PASSWORD in .env)"
            msg = MIMEMultipart()
            msg['From']    = username
            msg['To']      = to_email
            msg['Subject'] = subject
            if html_body:
                msg.attach(MIMEText(body, 'plain'))
                msg.attach(MIMEText(html_body, 'html'))
            else:
                msg.attach(MIMEText(body, 'plain'))
            if attachments:
                for att in attachments:
                    path = att if os.path.exists(att) else os.path.join(OUTPUT_DIR, att)
                    if os.path.exists(path):
                        with open(path, 'rb') as f:
                            part = MIMEApplication(f.read(), Name=os.path.basename(path))
                        part['Content-Disposition'] = f'attachment; filename="{os.path.basename(path)}"'
                        msg.attach(part)
            server = smtplib.SMTP(EMAIL_CONFIG.get('smtp_server', 'smtp.gmail.com'),
                                  EMAIL_CONFIG.get('smtp_port', 587))
            server.starttls()
            server.login(username, password)
            server.send_message(msg)
            server.quit()
            return f"Email sent to {to_email}"
        except Exception as e:
            return f"Error: {str(e)}"

# ==================== TOOL BOX ====================
class ToolBox:
    @staticmethod
    def create_pdf(filename: str, content, title: str = "") -> str:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.lib.enums import TA_CENTER
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from xml.sax.saxutils import escape
            filename = sanitize_filename(filename)
            filepath = os.path.join(OUTPUT_DIR, filename)
            if isinstance(content, list):
                content = "\n\n".join(str(i) for i in content)
            elif not isinstance(content, str):
                content = str(content)
            font_name = 'Helvetica'
            try:
                r = subprocess.run(['fc-match','--format=%{file}','DejaVuSans'],
                                   capture_output=True, text=True, timeout=2)
                dp = r.stdout.strip()
                if dp and os.path.exists(dp):
                    pdfmetrics.registerFont(TTFont('DejaVu', dp))
                    font_name = 'DejaVu'
            except Exception:
                pass
            doc = SimpleDocTemplate(filepath, pagesize=A4,
                                    topMargin=0.75*inch, bottomMargin=0.75*inch,
                                    leftMargin=0.75*inch, rightMargin=0.75*inch)
            styles = getSampleStyleSheet()
            story = []
            if title:
                ts = ParagraphStyle('T', parent=styles['Heading1'],
                                    fontName=font_name, fontSize=18,
                                    alignment=TA_CENTER, spaceAfter=20)
                story.append(Paragraph(escape(title), ts))
                story.append(Spacer(1, 0.3*inch))
            bs = ParagraphStyle('B', parent=styles['Normal'],
                                fontName=font_name, fontSize=12, spaceAfter=12)
            for para in content.split('\n\n'):
                for line in para.split('\n'):
                    if line.strip():
                        story.append(Paragraph(escape(line.strip()), bs))
                        story.append(Spacer(1, 0.1*inch))
            doc.build(story)
            return f"PDF created: {filepath}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def create_docx(filename: str, content: str, title: str = "") -> str:
        try:
            filename = sanitize_filename(filename)
            doc = Document()
            if title:
                doc.add_heading(title, 0)
            for para in content.split('\n\n'):
                if para.strip():
                    doc.add_paragraph(para.strip())
            filepath = os.path.join(OUTPUT_DIR, filename)
            doc.save(filepath)
            return f"DOCX created: {filepath}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def create_chart(filename: str, chart_type: str, data: Dict, title: str = "") -> str:
        try:
            filename = sanitize_filename(filename)
            filepath = os.path.join(OUTPUT_DIR, filename)
            plt.figure(figsize=(10, 6))
            if chart_type == "bar":
                plt.bar(data['labels'], data['values'])
            elif chart_type == "line":
                plt.plot(data['labels'], data['values'], marker='o')
            elif chart_type == "pie":
                plt.pie(data['values'], labels=data['labels'], autopct='%1.1f%%')
            else:
                return "Error: Use 'bar', 'line', or 'pie'."
            if title:
                plt.title(title)
            plt.tight_layout()
            plt.savefig(filepath, dpi=300, bbox_inches='tight')
            plt.close()
            return f"Chart created: {filepath}"
        except Exception as e:
            try: plt.close()
            except Exception: pass
            return f"Error: {str(e)}"

    @staticmethod
    def create_qr(data: str, filename: str = "qrcode.png") -> str:
        try:
            filename = sanitize_filename(filename)
            qr = qrcode.QRCode(version=1, box_size=10, border=5)
            qr.add_data(data)
            qr.make(fit=True)
            img = qr.make_image(fill_color="black", back_color="white")
            filepath = os.path.join(OUTPUT_DIR, filename)
            img.save(filepath)
            return f"QR code: {filepath}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def calculate(expression: str) -> str:
        return safe_calculate(expression)

    @staticmethod
    def translate_text(text: str, source_lang: str = "auto", target_lang: str = "ro") -> str:
        try:
            translated = GoogleTranslator(source=source_lang, target=target_lang).translate(text)
            return f"Translation ({source_lang}→{target_lang}): {translated}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    async def scrape_website(url: str, selector: str = None) -> str:
        if not is_safe_url(url):
            return "Error: URL not allowed (must be public http/https)"
        try:
            import httpx
            async with httpx.AsyncClient(timeout=10, follow_redirects=False) as client:
                response = await client.get(url)
                soup = BeautifulSoup(response.content, 'html.parser')
                if selector:
                    text = "\n".join(e.get_text().strip() for e in soup.select(selector))
                else:
                    text = soup.get_text()
                return f"Scraped:\n{text[:1000]}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def resize_image(input_path: str, output_filename: str, width: int, height: int = None) -> str:
        try:
            output_filename = sanitize_filename(output_filename)
            with Image.open(input_path) as img:
                if height is None:
                    height = int(img.height * width / img.width)
                img_resized = img.resize((width, height), Image.Resampling.LANCZOS)
                filepath = os.path.join(OUTPUT_DIR, output_filename)
                img_resized.save(filepath)
            return f"Image resized: {filepath}"
        except Exception as e:
            return f"Error: {str(e)}"

    @staticmethod
    def list_files() -> str:
        try:
            files = os.listdir(OUTPUT_DIR)
            if not files:
                return "No files"
            pdf   = sorted(f for f in files if f.endswith('.pdf'))[:10]
            docs  = sorted(f for f in files if f.endswith(('.docx','.doc')))[:10]
            imgs  = sorted(f for f in files if f.lower().endswith(('.png','.jpg','.jpeg','.gif')))[:10]
            other = sorted(f for f in files if f not in pdf + docs + imgs)[:10]
            result = "Files:\n"
            if pdf:   result += f"PDFs: {', '.join(pdf)}\n"
            if docs:  result += f"Docs: {', '.join(docs)}\n"
            if imgs:  result += f"Images: {', '.join(imgs)}\n"
            if other: result += f"Other: {', '.join(other)}"
            return result
        except Exception as e:
            return f"Error: {str(e)}"

# ==================== TOOL REGISTRY ====================
document_tools = DocumentTools()
email_tools    = EmailTools()

TOOLS = {
    "create_pdf":              ToolBox.create_pdf,
    "create_docx":             ToolBox.create_docx,
    "create_chart":            ToolBox.create_chart,
    "create_qr":               ToolBox.create_qr,
    "calculate":               ToolBox.calculate,
    "translate_text":          ToolBox.translate_text,
    "scrape_website":          ToolBox.scrape_website,
    "resize_image":            ToolBox.resize_image,
    "list_files":              ToolBox.list_files,
    "markdown_to_pdf":         document_tools.markdown_to_pdf,
    "markdown_to_html":        document_tools.markdown_to_html,
    "docx_to_pdf":             document_tools.docx_to_pdf,
    "merge_pdfs":              document_tools.merge_pdfs,
    "split_pdf":               document_tools.split_pdf,
    "compress_pdf":            document_tools.compress_pdf,
    "extract_images_from_pdf": document_tools.extract_images_from_pdf,
    "send_email":              email_tools.send_email,
}

# ==================== NATIVE TOOL SCHEMAS (Ollama/OpenAI format) ====================
TOOL_SCHEMAS = [
    {"type":"function","function":{"name":"create_pdf","description":"Create a PDF document","parameters":{"type":"object","properties":{"filename":{"type":"string","description":"Output filename e.g. report.pdf"},"content":{"type":"string","description":"Text content"},"title":{"type":"string","description":"Optional heading title","default":""}},"required":["filename","content"]}}},
    {"type":"function","function":{"name":"create_docx","description":"Create a Word .docx document","parameters":{"type":"object","properties":{"filename":{"type":"string"},"content":{"type":"string"},"title":{"type":"string","default":""}},"required":["filename","content"]}}},
    {"type":"function","function":{"name":"create_chart","description":"Create a bar, line, or pie chart as PNG","parameters":{"type":"object","properties":{"filename":{"type":"string"},"chart_type":{"type":"string","enum":["bar","line","pie"]},"data":{"type":"object","properties":{"labels":{"type":"array","items":{"type":"string"}},"values":{"type":"array","items":{"type":"number"}}},"required":["labels","values"]},"title":{"type":"string","default":""}},"required":["filename","chart_type","data"]}}},
    {"type":"function","function":{"name":"create_qr","description":"Generate a QR code image","parameters":{"type":"object","properties":{"data":{"type":"string","description":"URL or text to encode"},"filename":{"type":"string","default":"qrcode.png"}},"required":["data"]}}},
    {"type":"function","function":{"name":"calculate","description":"Evaluate a math expression safely","parameters":{"type":"object","properties":{"expression":{"type":"string","description":"e.g. 15*24 or (100+50)/3"}},"required":["expression"]}}},
    {"type":"function","function":{"name":"translate_text","description":"Translate text to another language","parameters":{"type":"object","properties":{"text":{"type":"string"},"source_lang":{"type":"string","default":"auto"},"target_lang":{"type":"string","description":"Language code e.g. ro, fr, de, es, en"}},"required":["text","target_lang"]}}},
    {"type":"function","function":{"name":"scrape_website","description":"Scrape text from a public website URL","parameters":{"type":"object","properties":{"url":{"type":"string"},"selector":{"type":"string","description":"Optional CSS selector"}},"required":["url"]}}},
    {"type":"function","function":{"name":"resize_image","description":"Resize an image file","parameters":{"type":"object","properties":{"input_path":{"type":"string"},"output_filename":{"type":"string"},"width":{"type":"integer"},"height":{"type":"integer","description":"Optional, preserves ratio if omitted"}},"required":["input_path","output_filename","width"]}}},
    {"type":"function","function":{"name":"list_files","description":"List all files in the output directory","parameters":{"type":"object","properties":{}}}},
    {"type":"function","function":{"name":"merge_pdfs","description":"Merge multiple PDF files into one","parameters":{"type":"object","properties":{"pdf_files":{"type":"array","items":{"type":"string"}},"output_filename":{"type":"string"}},"required":["pdf_files","output_filename"]}}},
    {"type":"function","function":{"name":"split_pdf","description":"Split a PDF into pages or extract specific pages","parameters":{"type":"object","properties":{"pdf_path":{"type":"string"},"pages":{"type":"array","items":{"type":"integer"},"description":"Page numbers to extract, or omit for all pages"},"output_prefix":{"type":"string","default":"split"}},"required":["pdf_path"]}}},
    {"type":"function","function":{"name":"compress_pdf","description":"Compress a PDF to reduce file size","parameters":{"type":"object","properties":{"input_path":{"type":"string"},"output_filename":{"type":"string"}},"required":["input_path","output_filename"]}}},
    {"type":"function","function":{"name":"markdown_to_pdf","description":"Convert Markdown text to PDF","parameters":{"type":"object","properties":{"markdown_content":{"type":"string"},"filename":{"type":"string"}},"required":["markdown_content","filename"]}}},
    {"type":"function","function":{"name":"markdown_to_html","description":"Convert Markdown text to HTML file","parameters":{"type":"object","properties":{"markdown_content":{"type":"string"},"filename":{"type":"string"}},"required":["markdown_content","filename"]}}},
    {"type":"function","function":{"name":"send_email","description":"Send an email with optional file attachments","parameters":{"type":"object","properties":{"to_email":{"type":"string"},"subject":{"type":"string"},"body":{"type":"string"},"attachments":{"type":"array","items":{"type":"string"},"description":"Filenames in output dir to attach"}},"required":["to_email","subject","body"]}}},
]

# Core subset for 4B model — small models are unreliable with >8 tools
CORE_TOOL_SCHEMAS = [
    s for s in TOOL_SCHEMAS
    if s["function"]["name"] in {
        "create_pdf", "create_docx", "create_chart", "create_qr",
        "calculate", "translate_text", "list_files", "send_email"
    }
]

# ==================== TOOL EXECUTION ====================
async def execute_tool(tool_name: str, params: Dict) -> str:
    """Execute a tool by name with given params, handling sync/async transparently"""
    try:
        import inspect
        func = TOOLS.get(tool_name)
        if not func:
            return f"Error: unknown tool '{tool_name}'"
        sig = inspect.signature(func)
        filtered = {k: v for k, v in params.items() if k in sig.parameters}
        if inspect.iscoroutinefunction(func):
            return await func(**filtered)
        return await asyncio.get_running_loop().run_in_executor(executor, lambda: func(**filtered))
    except Exception as e:
        logger.error(f"Tool {tool_name} failed: {e}")
        return f"Error executing {tool_name}: {str(e)}"

# ==================== OLLAMA INFERENCE WITH NATIVE TOOL LOOP ====================
async def call_ollama(model: str, messages: list, thinking: bool = False,
                      use_tools: bool = False) -> tuple[str, list]:
    msgs = list(messages)

    if not thinking and msgs and msgs[0]["role"] == "system":
        if "/no-think" not in msgs[0]["content"]:
            msgs[0] = dict(msgs[0])
            msgs[0]["content"] += "\n/no-think"

    tool_summary = []

    for loop in range(MAX_TOOL_LOOPS + 1):
        payload = {
            "model": model,
            "messages": msgs,
            "stream": False,
            "think": thinking,
            "options": {
                "temperature": 0.1 if use_tools else 0.4,
                "num_ctx": 4096 if use_tools else 8192,
                "num_thread": N_TOOL_THREADS,
            }
        }

        if use_tools:
            payload["tools"] = CORE_TOOL_SCHEMAS

        _timeout = aiohttp.ClientTimeout(total=None, sock_connect=60, sock_read=600)
        try:
            async with aiohttp.ClientSession() as session:
                async with session.post(
                    f"{OLLAMA_URL}/api/chat",
                    json=payload,
                    timeout=_timeout
                ) as resp:
                    if resp.status != 200:
                        err = await resp.text()
                        raise RuntimeError(f"Ollama error {resp.status}: {err}")
                    data = await resp.json()
        except Exception as e:
            logger.error(f"Ollama call failed ({model}): {e}")
            raise

        msg        = data.get("message", {})
        tool_calls = msg.get("tool_calls", [])

        # No tool calls or loop exhausted → return text
        # NOTE: done_reason intentionally ignored — Qwen often returns "stop" even with tool_calls
        if not tool_calls or loop >= MAX_TOOL_LOOPS:
            text = msg.get("content", "").strip()
            text = re.sub(r'<think>.*?</think>', '', text, flags=re.DOTALL).strip()
            return text, tool_summary

        msgs.append({
            "role": "assistant",
            "content": msg.get("content", ""),
            "tool_calls": tool_calls
        })

        for tc in tool_calls:
            func_info  = tc.get("function", {})
            tool_name  = func_info.get("name", "")
            tool_args  = func_info.get("arguments", {})
            if isinstance(tool_args, str):
                try:
                    tool_args = json.loads(tool_args)
                except Exception:
                    tool_args = {}

            logger.info(f"[Tool] {tool_name}({tool_args})")
            result = await execute_tool(tool_name, tool_args)
            logger.info(f"[Tool result] {result[:120]}")
            tool_summary.append(f"[{tool_name}] {result}")

            msgs.append({
                "role": "tool",
                "tool_name": tool_name,
                "content": result
            })

    return "Operations completed.", tool_summary

# ==================== WEB SEARCH ====================
searxng_available = False

async def check_searxng():
    global searxng_available
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(f"{SEARXNG_URL}/search?q=test&format=json",
                                   timeout=aiohttp.ClientTimeout(total=5)) as resp:
                searxng_available = resp.status == 200
                logger.info(f"SearXNG: {'available' if searxng_available else 'unavailable'}")
    except Exception:
        searxng_available = False
        logger.warning("SearXNG unavailable")

async def perform_web_search(query: str, num_results: int = 3) -> List[Dict]:
    if not searxng_available:
        return []
    try:
        params = {"q": query, "format": "json", "language": "en", "categories": "general"}
        async with aiohttp.ClientSession() as session:
            async with session.get(f"{SEARXNG_URL}/search", params=params,
                                   timeout=aiohttp.ClientTimeout(total=8)) as resp:
                if resp.status == 200:
                    data = await resp.json()
                    return [{"title": r.get("title",""),
                             "content": r.get("content","")[:300],
                             "url": r.get("url","")}
                            for r in data.get("results",[])[:num_results]]
    except Exception as e:
        logger.error(f"Search failed: {e}")
    return []

async def check_ollama() -> bool:
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(f"{OLLAMA_URL}/api/tags",
                                   timeout=aiohttp.ClientTimeout(total=5)) as resp:
                if resp.status == 200:
                    data = await resp.json()
                    available = [m["name"] for m in data.get("models", [])]
                    logger.info(f"Ollama models available: {available}")
                    for model in [TOOL_MODEL, CHAT_MODEL]:
                        if not any(model in a for a in available):
                            logger.warning(f"Model {model} not found in Ollama")
                    return True
    except Exception as e:
        logger.error(f"Ollama not reachable: {e}")
    return False

# ==================== REQUEST DEDUPLICATION ====================
class RequestCache:
    def __init__(self, max_size=100):
        self.cache = OrderedDict()
        self.lock = Lock()
        self.max_size = max_size

    def is_duplicate(self, request_hash: str, window_seconds: int = 5) -> bool:
        with self.lock:
            if request_hash in self.cache:
                if time.time() - self.cache[request_hash] < window_seconds:
                    return True
            self.cache[request_hash] = time.time()
            if len(self.cache) > self.max_size:
                self.cache.popitem(last=False)
            return False

request_cache = RequestCache()

def is_openwebui_search_request(msg: str) -> bool:
    indicators = ["analyze the chat history","generate search queries",
                  "### task:","necessity of generating search queries"]
    return any(i in msg.lower() for i in indicators)

def should_use_tool_model(user_msg: str) -> bool:
    lower = user_msg.lower()
    return any(kw in lower for kw in TOOL_KEYWORDS)

# ==================== SYSTEM PROMPTS ====================
SYSTEM_PROMPT = {
    "en": "You are a helpful AI assistant. When the user asks you to create, calculate, translate, or perform any action, call the appropriate tool. For normal conversation, respond normally.",
    "ro": "Ești un asistent AI util. Când utilizatorul cere să creezi, calculezi, traduci sau efectuezi acțiuni, folosește instrumentul potrivit. Pentru conversații normale, răspunde normal.",
    "fr": "Vous êtes un assistant IA utile. Quand l'utilisateur demande de créer, calculer, traduire ou effectuer des actions, appelez l'outil approprié. Pour les conversations normales, répondez normalement.",
}

ROUTER_PROMPT = """Classify the user message. Reply with exactly one word:
TOOL — if the user wants to create a file, calculate, translate, scrape, generate, or perform an action
CHAT — if the user wants conversation, explanation, or information

Examples:
"Hello" -> CHAT
"Create a PDF" -> TOOL
"Calculate 5*7" -> TOOL
"What is Python?" -> CHAT
"Make a chart" -> TOOL
"Thanks" -> CHAT
/no-think"""

def build_system_message(lang: str, web_results: List[Dict] = None) -> str:
    msg = SYSTEM_PROMPT.get(lang, SYSTEM_PROMPT["en"])
    if web_results:
        msg += "\n\n=== SEARCH RESULTS ===\n"
        for i, r in enumerate(web_results[:3], 1):
            msg += f"[{i}] {r['title']}\n{r['content'][:200]}\n\n"
    return msg

# ==================== API MODELS ====================
class Message(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    model: str
    messages: List[Message]
    temperature: Optional[float] = 0.4
    max_tokens: Optional[int] = 1024
    stream: Optional[bool] = False
    web_search: Optional[bool] = False

# ==================== LIFESPAN ====================
@asynccontextmanager
async def lifespan(app: FastAPI):
    logger.info("=" * 50)
    logger.info("AI AGENT SERVER STARTING (Dual Model + Native Tool Calling)")
    logger.info("=" * 50)
    validate_environment()
    ollama_ok = await check_ollama()
    if not ollama_ok:
        raise RuntimeError("Ollama is not reachable — is it running?")
    await check_searxng()
    logger.info(f"Tool model:  {TOOL_MODEL}")
    logger.info(f"Chat model:  {CHAT_MODEL}")
    logger.info(f"Tools:       {len(TOOLS)}")
    logger.info(f"Web search:  {searxng_available}")
    logger.info("=" * 50)
    yield
    logger.info("Shutting down...")
    executor.shutdown(wait=True)
    logger.info("Shutdown complete")

# ==================== FASTAPI APP ====================
app = FastAPI(title="AI Agent - Native Tool Calling", lifespan=lifespan)

limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

allowed_origins = os.getenv("ALLOWED_ORIGINS", "http://localhost:3000").split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
    max_age=3600,
)

@app.middleware("http")
async def limit_request_size(request: Request, call_next):
    if request.method == "POST":
        cl = request.headers.get("content-length")
        if cl and int(cl) > 1024 * 1024:
            return JSONResponse(status_code=413, content={"error": "Request too large"})
    return await call_next(request)

# ==================== ENDPOINTS ====================
@app.get("/")
async def root():
    return {"status": "running", "tool_model": TOOL_MODEL, "chat_model": CHAT_MODEL,
            "tools": len(TOOLS), "web_search": searxng_available,
            "mode": "native_tool_calling"}

@app.get("/v1/models")
async def list_models():
    return {"object": "list", "data": [
        {"id": CHAT_MODEL, "object": "model", "created": int(time.time()), "owned_by": "ollama"},
        {"id": TOOL_MODEL, "object": "model", "created": int(time.time()), "owned_by": "ollama"},
    ]}

@app.get("/api/tags")
async def list_models_ollama():
    return {"models": [
        {"name": CHAT_MODEL, "model": CHAT_MODEL, "modified_at": "2025-01-01T00:00:00Z"},
        {"name": TOOL_MODEL, "model": TOOL_MODEL, "modified_at": "2025-01-01T00:00:00Z"},
    ]}

@app.post("/v1/chat/completions")
@limiter.limit("10/minute")
async def chat_completion(request: Request, body: ChatRequest):
    try:
        messages  = [m.model_dump() for m in body.messages]
        last_user = next((m["content"] for m in reversed(messages) if m["role"] == "user"), "")
        is_owui   = is_openwebui_search_request(last_user)

        req_hash = hashlib.md5(f"{last_user}_{body.model}".encode()).hexdigest()
        if not is_owui and request_cache.is_duplicate(req_hash):
            raise HTTPException(status_code=429, detail="Duplicate request, please wait")

        lang = detect_language(last_user)

        web_results = []
        if body.web_search and searxng_available and last_user and not is_owui:
            try:
                web_results = await asyncio.wait_for(
                    perform_web_search(last_user), timeout=5.0)
            except asyncio.TimeoutError:
                logger.warning("Web search timed out")

        chat_messages = [{"role": "system", "content": build_system_message(lang, web_results)}]
        chat_messages.extend(messages[-6:])

        # ── Model routing ──────────────────────────────────────────────────────
        explicit_model = body.model if body.model in (TOOL_MODEL, CHAT_MODEL) else None
        if explicit_model:
            selected_model = explicit_model
            use_tool_model = (selected_model == TOOL_MODEL)
            logger.info(f"Explicit model: {selected_model}")
        elif is_owui:
            use_tool_model = False
            selected_model = CHAT_MODEL
        else:
            try:
                router_msgs = [
                    {"role": "system", "content": ROUTER_PROMPT},
                    {"role": "user", "content": last_user}
                ]
                router_payload = {
                    "model": TOOL_MODEL,
                    "messages": router_msgs,
                    "stream": False,
                    "think": False,
                    "options": {"temperature": 0.0, "num_ctx": 512, "num_thread": N_TOOL_THREADS}
                }
                _rt = aiohttp.ClientTimeout(total=None, sock_connect=10, sock_read=30)
                async with aiohttp.ClientSession() as _s:
                    async with _s.post(f"{OLLAMA_URL}/api/chat", json=router_payload, timeout=_rt) as _r:
                        _rd = await _r.json()
                        router_decision = _rd.get("message", {}).get("content", "CHAT").strip().upper()
                use_tool_model = "TOOL" in router_decision
                logger.info(f"Router decision: {router_decision} → tool_model={use_tool_model}")
            except Exception as re:
                logger.warning(f"Router failed, falling back to keyword match: {re}")
                use_tool_model = should_use_tool_model(last_user)

            selected_model = TOOL_MODEL if use_tool_model else CHAT_MODEL
            logger.info(f"Auto-routed to: {selected_model}")

        enable_thinking = (not use_tool_model) and len(last_user.split()) > 10
        use_tools = use_tool_model and not is_owui

        logger.info(f"User: {last_user[:80]}")
        logger.info(f"Model: {selected_model} | thinking={enable_thinking} | tools={use_tools}")

        content, tool_results = await call_ollama(
            selected_model, chat_messages,
            thinking=enable_thinking,
            use_tools=use_tools
        )

        if not content and tool_results:
            content = "\n\n".join(tool_results)

        logger.info(f"Response: {content[:100]}")

        return {
            "id": f"chatcmpl-{uuid.uuid4()}",
            "object": "chat.completion",
            "created": int(time.time()),
            "model": selected_model,
            "choices": [{"index": 0,
                         "message": {"role": "assistant", "content": content},
                         "finish_reason": "stop"}],
            "usage": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0}
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Chat error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/files/{filename}")
async def download_file(filename: str):
    filename = sanitize_filename(filename)
    filepath = os.path.join(OUTPUT_DIR, filename)
    if os.path.exists(filepath):
        return FileResponse(filepath)
    raise HTTPException(status_code=404, detail="File not found")

@app.get("/api/list-files")
async def api_list_files():
    try:
        files = sorted(os.listdir(OUTPUT_DIR))
        return {
            "files": files,
            "categorized": {
                "pdfs":   [f for f in files if f.endswith('.pdf')],
                "docs":   [f for f in files if f.endswith(('.docx', '.doc'))],
                "images": [f for f in files if f.lower().endswith(('.png', '.jpg', '.jpeg', '.gif'))],
                "other":  [f for f in files if not f.endswith(('.pdf', '.docx', '.doc', '.png', '.jpg', '.jpeg', '.gif'))]
            },
            "total": len(files),
            "download_base": "/files/"
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/tools")
async def list_tools_endpoint():
    return {"tools": list(TOOLS.keys()), "total": len(TOOLS)}

@app.get("/health")
async def health():
    ollama_ok = await check_ollama()
    return {"status": "healthy" if ollama_ok else "degraded",
            "ollama": ollama_ok, "tool_model": TOOL_MODEL, "chat_model": CHAT_MODEL,
            "tools": len(TOOLS), "web_search": searxng_available,
            "timestamp": datetime.utcnow().isoformat()}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000, log_level="info")
